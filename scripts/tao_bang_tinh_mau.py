"""Tạo lại file data/bo-cot-thuoc-tinh-thoi-trang.xlsx (bảng tính mẫu). Chạy: python scripts/tao_bang_tinh_mau.py"""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ENUMS = {
 "gioi_tinh": ["nữ","nam","unisex"],
 "danh_muc": ["áo","đầm","chân váy","quần","set đồ","áo khoác","đồ ngủ","đồ bơi"],
 "kieu_sp": ["áo thun","áo sơ mi","áo kiểu","áo croptop","áo hai dây","áo len","áo polo",
             "đầm suông","đầm ôm body","đầm xòe","đầm babydoll","đầm maxi","đầm sơ mi",
             "chân váy chữ A","chân váy bút chì","chân váy xếp ly","chân váy jean",
             "quần jean","quần tây","quần ống rộng","quần short","quần jogger","quần legging",
             "áo khoác jean","áo blazer","áo cardigan","áo khoác gió","áo phao","khác"],
 "kieu_co": ["trễ vai","lệch vai","cổ V","cổ tròn","cổ vuông","cổ yếm","hai dây","cúp ngực",
             "cổ sơ mi","cổ đức/bẻ","cổ lọ","cổ thuyền","cổ tim","không áp dụng"],
 "tay_ao": ["sát nách","tay ngắn","tay lỡ","tay dài","tay phồng","tay loe","không áp dụng"],
 "do_dai": ["croptop","ngang hông","qua mông","trên gối","ngang gối","midi","maxi","không áp dụng"],
 "dang": ["ôm","suông","rộng/oversize","xòe","chữ A","babydoll","đứng"],
 "chat_lieu": ["cotton","thun","lụa/satin","voan/chiffon","linen/đũi","jean","kaki","len","nỉ","ren","da","polyester","khác"],
 "hoa_tiet": ["trơn","hoa nhí","hoa to","kẻ sọc","caro","chấm bi","in chữ/hình","ren","loang màu"],
 "mau_chinh": ["trắng","đen","be/kem","nâu","xám","hồng","đỏ","cam","vàng","xanh lá","xanh dương","xanh navy","tím","nhiều màu"],
 "phong_cach": ["basic","nữ tính/tiểu thư","công sở","cá tính/streetwear","vintage","quyến rũ","thể thao","hàn quốc"],
 "mua": ["xuân hè","thu đông","quanh năm"],
 "co_gian": ["không","ít","nhiều"],
 "trang_thai": ["đang bán","hết hàng","ngừng bán"],
 "size": ["XS","S","M","L","XL","XXL","Freesize"],
}
MULTI = {
 "dip": ["đi làm","đi học","dạo phố","hẹn hò","đi biển","dự tiệc","ở nhà","tập thể thao"],
 "chi_tiet": ["bèo nhún","nơ","xẻ tà","cut-out","đính cúc","thắt eo/đai","có túi","có lót trong","phối ren","dây rút","khóa kéo"],
}

PRODUCT_COLS = [
 ("ma_sp","Mã sản phẩm, duy nhất (VD: SP0001)",None),
 ("ten_sp","Tên hiển thị cho khách",None),
 ("anh_chinh_url","Link ảnh chính",None),
 ("anh_phu_url","Link ảnh phụ, cách nhau dấu phẩy",None),
 ("gia_ban","Giá bán, VND",None),
 ("gia_goc","Giá gốc trước giảm, VND (để trống nếu không giảm)",None),
 ("gioi_tinh","",  "gioi_tinh"),
 ("danh_muc","Nhóm lớn", "danh_muc"),
 ("kieu_sp","Kiểu cụ thể", "kieu_sp"),
 ("kieu_co","Kiểu cổ / vai", "kieu_co"),
 ("tay_ao","", "tay_ao"),
 ("do_dai","Độ dài thân áo / váy / quần", "do_dai"),
 ("dang","Form dáng", "dang"),
 ("chat_lieu","", "chat_lieu"),
 ("hoa_tiet","", "hoa_tiet"),
 ("mau_chinh","Màu chủ đạo trên ảnh chính", "mau_chinh"),
 ("phong_cach","", "phong_cach"),
 ("dip","Nhiều giá trị, cách nhau dấu phẩy (xem sheet DanhMuc)", None),
 ("mua","", "mua"),
 ("chi_tiet","Nhiều giá trị, cách nhau dấu phẩy (xem sheet DanhMuc)", None),
 ("co_gian","Độ co giãn vải", "co_gian"),
 ("mo_ta_tu_do","2-3 câu mô tả kiểu dáng, dùng để tạo embedding", None),
 ("trang_thai","", "trang_thai"),
 ("gan_boi","AI hay người gắn (ai / nguoi / ai_da_duyet)", None),
]
SAMPLES = [
 ["SP0001","Áo trễ vai bèo nhún Hạ Vy","https://…/sp0001.jpg","",289000,349000,"nữ","áo","áo kiểu","trễ vai","tay ngắn","croptop","rộng/oversize","voan/chiffon","trơn","trắng","nữ tính/tiểu thư","đi biển, dạo phố, hẹn hò","xuân hè","bèo nhún, có lót trong","ít","Áo trễ vai voan trắng, viền bèo nhún quanh vai, dáng croptop rộng nhẹ, hợp đi biển và chụp ảnh.","đang bán","ai_da_duyet"],
 ["SP0002","Đầm maxi hoa nhí cổ vuông","https://…/sp0002.jpg","",459000,"","nữ","đầm","đầm maxi","cổ vuông","tay phồng","maxi","xòe","voan/chiffon","hoa nhí","vàng","vintage","đi biển, dạo phố","xuân hè","thắt eo/đai, xẻ tà","không","Đầm maxi nền vàng hoa nhí, cổ vuông tay phồng, chun eo, xẻ tà nhẹ, phong cách vintage.","đang bán","ai"],
 ["SP0003","Quần ống rộng công sở","https://…/sp0003.jpg","",329000,"","nữ","quần","quần ống rộng","không áp dụng","không áp dụng","maxi","rộng/oversize","polyester","trơn","be/kem","công sở","đi làm","quanh năm","có túi, khóa kéo","ít","Quần ống rộng lưng cao màu be, vải đứng form, có túi chéo, hợp đi làm.","đang bán","nguoi"],
]

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
NOTE_FONT = Font(italic=True, color="666666", size=9)

wb = Workbook()
# --- Huong dan
ws = wb.active; ws.title = "HuongDan"
lines = [
 ("Bộ cột thuộc tính cho shop thời trang", Font(bold=True, size=14)),
 ("",None),
 ("SanPham: mỗi dòng một mẫu sản phẩm. Cột có mũi tên chọn giá trị từ danh sách cố định, để bot lọc chính xác.",None),
 ("BienThe: mỗi dòng một tổ hợp màu + size của một mẫu, kèm tồn kho. Bot chỉ gợi ý biến thể còn hàng.",None),
 ("BangSize: bảng quy đổi để bot tư vấn size theo cân nặng, chiều cao, số đo.",None),
 ("DanhMuc: các giá trị cho phép. Thêm giá trị mới ở đây (và trong schema JSON) chứ không gõ tự do trong SanPham.",None),
 ("",None),
 ("Cột dip và chi_tiet nhận nhiều giá trị, cách nhau dấu phẩy, chỉ dùng giá trị có trong DanhMuc.",None),
 ("Cột mo_ta_tu_do không dùng để lọc; nó được đưa vào embedding để tìm theo nghĩa.",None),
 ("Cột gan_boi giúp biết dòng nào do AI gắn mà người chưa kiểm tra.",None),
 ("",None),
 ("Prompt gợi ý cho Gemini (kèm ảnh + schema trong file gemini-schema-thuoc-tinh.json):",Font(bold=True)),
 ("Bạn là nhân viên nhập liệu của shop thời trang. Nhìn ảnh sản phẩm và điền thuộc tính theo schema. Chỉ chọn giá trị có trong danh sách. Nếu không nhìn thấy rõ thì chọn giá trị gần nhất, với kiểu cổ/tay không liên quan (VD quần) chọn \"không áp dụng\". mo_ta_tu_do viết 2-3 câu tiếng Việt tả kiểu dáng, màu, chất liệu, dịp mặc.",None),
 ("",None),
 ("Nguồn bộ cột: dựa trên các thuộc tính sản phẩm mà người bán thời trang hay phải điền trên sàn và các bộ lọc thời trang phổ biến; chưa đối chiếu trực tiếp với trang người bán của Shopee.",NOTE_FONT),
]
for i,(t,f) in enumerate(lines,1):
    c = ws.cell(row=i, column=1, value=t)
    if f: c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 120

# --- DanhMuc
dm = wb.create_sheet("DanhMuc")
ranges = {}
all_lists = {**ENUMS, **MULTI}
for ci,(k,vals) in enumerate(all_lists.items(),1):
    col = get_column_letter(ci)
    h = dm.cell(row=1,column=ci,value=k + (" (nhiều)" if k in MULTI else ""))
    h.fill=HEAD_FILL; h.font=HEAD_FONT
    for ri,v in enumerate(vals,2):
        dm.cell(row=ri,column=ci,value=v)
    ranges[k] = f"DanhMuc!${col}$2:${col}${len(vals)+1}"
    dm.column_dimensions[col].width = max(14, max(len(v) for v in vals)+2)
dm.freeze_panes="A2"

def header(ws, cols, notes=True):
    for ci,(name,note,_) in enumerate(cols,1):
        c=ws.cell(row=1,column=ci,value=name); c.fill=HEAD_FILL; c.font=HEAD_FONT
        if notes:
            n=ws.cell(row=2,column=ci,value=note); n.font=NOTE_FONT; n.alignment=Alignment(wrap_text=True)
        ws.column_dimensions[get_column_letter(ci)].width = 16
    ws.freeze_panes = "B3" if notes else "B2"

# --- SanPham
sp = wb.create_sheet("SanPham", 1)
header(sp, PRODUCT_COLS)
for ri,row in enumerate(SAMPLES,3):
    for ci,v in enumerate(row,1):
        sp.cell(row=ri,column=ci,value=v)
for ci,(name,_,enum) in enumerate(PRODUCT_COLS,1):
    col=get_column_letter(ci)
    if enum:
        dv=DataValidation(type="list", formula1="="+ranges[enum], allow_blank=True, showErrorMessage=True,
                          errorTitle="Giá trị không hợp lệ", error="Chọn giá trị trong sheet DanhMuc")
        sp.add_data_validation(dv); dv.add(f"{col}3:{col}2000")
for name,w in [("ten_sp",30),("mo_ta_tu_do",60),("dip",26),("chi_tiet",26),("anh_chinh_url",24)]:
    idx=[c[0] for c in PRODUCT_COLS].index(name)+1
    sp.column_dimensions[get_column_letter(idx)].width=w
for col in ("E","F"):
    for r in range(3,2001): sp[f"{col}{r}"].number_format='#,##0'

# --- BienThe
bt = wb.create_sheet("BienThe", 2)
BT_COLS=[("ma_bien_the","VD: SP0001-TRANG-M",None),("ma_sp","Khớp với SanPham",None),("mau","Tên màu khách thấy",None),
         ("size","",None),("ton_kho","Số lượng còn",None),("gia_rieng","Để trống nếu bằng gia_ban",None),("anh_url","Ảnh riêng của màu này (nếu có)",None)]
header(bt, BT_COLS)
rows=[["SP0001-TRANG-S","SP0001","trắng","S",12,"",""],["SP0001-TRANG-M","SP0001","trắng","M",8,"",""],
      ["SP0001-DEN-M","SP0001","đen","M",0,"",""],["SP0002-VANG-FREE","SP0002","vàng","Freesize",20,"",""],
      ["SP0003-BE-M","SP0003","be","M",5,"",""],["SP0003-BE-L","SP0003","be","L",3,"",""]]
for ri,row in enumerate(rows,3):
    for ci,v in enumerate(row,1): bt.cell(row=ri,column=ci,value=v)
dv=DataValidation(type="list", formula1="="+ranges["size"], allow_blank=True); bt.add_data_validation(dv); dv.add("D3:D5000")
dvi=DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0"); bt.add_data_validation(dvi); dvi.add("E3:E5000")
bt.column_dimensions["A"].width=22

# --- BangSize
bs = wb.create_sheet("BangSize", 3)
BS_COLS=[("size","",None),("can_nang_kg","Khoảng cân nặng",None),("chieu_cao_cm","",None),("vong_1_cm","Vòng ngực",None),("vong_2_cm","Vòng eo",None),("vong_3_cm","Vòng mông",None)]
header(bs, BS_COLS)
srows=[["S","40-47","148-155","78-82","60-64","84-88"],["M","48-53","153-160","82-86","64-68","88-92"],
       ["L","54-58","158-165","86-90","68-72","92-96"],["XL","59-64","160-168","90-95","72-77","96-100"]]
for ri,row in enumerate(srows,3):
    for ci,v in enumerate(row,1): bs.cell(row=ri,column=ci,value=v)
bs.cell(row=8,column=1,value="Số liệu mẫu tham khảo, mỗi shop tự sửa theo form của mình.").font=NOTE_FONT

import os; out=os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "bo-cot-thuoc-tinh-thoi-trang.xlsx")
wb.save(out)

# --- Gemini schema
props={}
for name,_,enum in PRODUCT_COLS:
    if enum and name not in ("trang_thai",):
        props[name]={"type":"STRING","enum":ENUMS[enum]}
for k,v in MULTI.items():
    props[k]={"type":"ARRAY","items":{"type":"STRING","enum":v}}
props["mo_ta_tu_do"]={"type":"STRING","description":"2-3 câu tiếng Việt tả kiểu dáng, màu, chất liệu, dịp mặc"}
schema={"type":"OBJECT","properties":props,"required":list(props.keys())}
pass
print(out, len(props))
