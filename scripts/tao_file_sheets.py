"""Tạo file Excel để tải lên Google Sheets: đúng mẫu data/bo-cot-thuoc-tinh-thoi-trang.xlsx (dòng 1 tên cột,
dòng 2 giải thích, dữ liệu từ dòng 3, giữ danh sách chọn và trang DanhMuc), điền sẵn 300 món demo bot đang bán.

Ảnh để dạng /anh/<ma_sp>.jpg (Worker tự phục vụ) và nội dung giữ nguyên như lúc nạp demo, nên lần đồng bộ đầu
dùng lại vector đã có, không tốn tiền tạo lại.

Đầu vào: data/bo-cot-thuoc-tinh-thoi-trang.xlsx, scripts/out/demo/*.csv (của tao_du_lieu_demo.py)
Đầu ra:  data/du-lieu-shop-demo.xlsx
Chạy:    python scripts/tao_file_sheets.py
"""
import csv
import os

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MAU = os.path.join(ROOT, "data", "bo-cot-thuoc-tinh-thoi-trang.xlsx")
DEMO = os.path.join(ROOT, "scripts", "out", "demo")
OUT = os.path.join(ROOT, "data", "du-lieu-shop-demo.xlsx")
SO = {"gia_ban", "gia_goc", "ton_kho", "gia_rieng"}


def doc(ten):
    with open(os.path.join(DEMO, ten), encoding="utf-8") as f:
        return list(csv.DictReader(f))


def dien(ws, rows, mac_dinh=None):
    """Xóa dòng mẫu từ dòng 3 rồi điền rows theo đúng thứ tự cột ở dòng 1 của trang."""
    cols = [c.value for c in ws[1] if c.value]
    for r in ws.iter_rows(min_row=3):
        for c in r:
            c.value = None
    for ri, row in enumerate(rows, 3):
        for ci, col in enumerate(cols, 1):
            v = row.get(col) or (mac_dinh or {}).get(col) or None
            if v is not None and col in SO:
                v = int(v)
            ws.cell(row=ri, column=ci, value=v)
    return len(rows)


wb = load_workbook(MAU)
n_sp = dien(wb["SanPham"], doc("san_pham.csv"), {"gan_boi": "ai"})
n_bt = dien(wb["BienThe"], doc("bien_the.csv"))
n_bs = dien(wb["BangSize"], doc("bang_size.csv"))

# ThongTinShop: bot trả lời câu hỏi về shop theo đúng trang này. Ô noi_dung để trống thì bot nói chưa có thông tin.
THONG_TIN = [
    ("Thanh toán", "Chỉ nhận thanh toán khi nhận hàng (COD)."),
    ("Phí ship", "Chưa gồm trong giá, khách trả cho bên giao hàng khi nhận; nhân viên shop báo phí cụ thể."),
    ("Thời gian giao hàng", "Khoảng 1-2 ngày sau khi đặt."),
    ("Đổi trả", ""),
    ("Địa chỉ cửa hàng", ""),
    ("Giờ làm việc", ""),
    ("Hotline", ""),
    ("Khuyến mãi đang chạy", ""),
]
tt = wb.create_sheet("ThongTinShop", 4)
for ci, (ten, ghi_chu, rong) in enumerate([("chu_de", "", 26), ("noi_dung", "Mỗi dòng một chủ đề. Để trống thì bot nói chưa có thông tin và nhờ nhân viên trả lời.", 90)], 1):
    h = tt.cell(row=1, column=ci, value=ten)
    h.fill, h.font = PatternFill("solid", fgColor="1F4E78"), Font(bold=True, color="FFFFFF")
    n = tt.cell(row=2, column=ci, value=ghi_chu or None)
    n.font, n.alignment = Font(italic=True, color="666666", size=9), Alignment(wrap_text=True)
    tt.column_dimensions["AB"[ci - 1]].width = rong
for ri, (chu_de, noi_dung) in enumerate(THONG_TIN, 3):
    tt.cell(row=ri, column=1, value=chu_de)
    tt.cell(row=ri, column=2, value=noi_dung or None).alignment = Alignment(wrap_text=True)
tt.freeze_panes = "A3"

# Trình bày giống hàm lamDepTrang_ trong apps-script/Bot.gs (script cũng tự làm lại lần đầu mở bảng tính).
RONG = {"ma_sp": 11, "ten_sp": 37, "anh_chinh_url": 23, "anh_phu_url": 20, "gia_ban": 14, "gia_goc": 14, "dip": 21,
        "chi_tiet": 21, "mo_ta_tu_do": 40, "trang_thai": 13, "gan_boi": 10, "ma_bien_the": 23, "size": 10,
        "ton_kho": 11, "gia_rieng": 14, "anh_url": 23}
TIEN = {"gia_ban", "gia_goc", "gia_rieng"}
GIUA = {"ma_sp", "size", "ton_kho", "trang_thai", "gan_boi"}
for ten in ["SanPham", "BienThe", "BangSize", "ThongTinShop"]:
    ws = wb[ten]
    ws.sheet_properties.tabColor = "3C78D8"
    ws.row_dimensions[1].height = 26
    for c in ws[1]:
        if not c.value:
            continue
        c.fill, c.font = PatternFill("solid", fgColor="1F3A5F"), Font(bold=True, color="FFFFFF", size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        g = ws.cell(row=2, column=c.column)
        g.fill = PatternFill("solid", fgColor="EEF3F8")
        g.font, g.alignment = Font(italic=True, color="5F6B7A", size=9), Alignment(vertical="top", wrap_text=True)
        if ten in ("SanPham", "BienThe") and c.value in RONG:
            ws.column_dimensions[c.column_letter].width = RONG[c.value]
        if c.value in TIEN or c.value in GIUA or ten == "BangSize":
            for (o,) in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=c.column, max_col=c.column):
                if c.value in TIEN:
                    o.number_format = '#,##0" ₫"'
                else:
                    o.alignment = Alignment(horizontal="center")
wb["SanPham"].freeze_panes = "C3"
for ten in ["HuongDan", "DanhMuc"]:
    if ten in wb.sheetnames:
        wb[ten].sheet_properties.tabColor = "999999"

wb.save(OUT)
print(f"{OUT}: {n_sp} sản phẩm, {n_bt} biến thể, {n_bs} dòng bảng size, {len(THONG_TIN)} chủ đề thông tin shop")
